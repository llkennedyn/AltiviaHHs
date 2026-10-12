"""ALTIVIA HH CONTROL - local management portal; run using INICIAR_ALTIVIA.bat."""
import os, sqlite3, secrets, io, calendar, json, unicodedata, re
from pathlib import Path
from datetime import datetime, date, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from flask import Flask, request, session, redirect, url_for, flash, render_template_string, abort, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from waitress import serve

ROOT=Path(__file__).resolve().parent
DB=ROOT/'altivia_hh.db'
SECRET=ROOT/'secrets.json'
if not SECRET.exists():
    SECRET.write_text(json.dumps({'secret':secrets.token_hex(32)}),encoding='utf8')
app=Flask(__name__)
app.secret_key=json.loads(SECRET.read_text(encoding='utf8'))['secret']
app.config.update(SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',SESSION_COOKIE_SECURE=False,MAX_CONTENT_LENGTH=1024*1024)

def conn():
    db=sqlite3.connect(DB,timeout=15)
    db.row_factory=sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.execute('PRAGMA busy_timeout=15000')
    return db

def init():
    # Migrate old one-liquidation-per-month constraint, preserving IDs and detail rows.
    raw=sqlite3.connect(DB)
    row=raw.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='liquidations'").fetchone()
    if row and 'UNIQUE(user_id,period)' in row[0].replace(' ',''):
        raw.execute('PRAGMA foreign_keys=OFF')
        raw.execute('BEGIN IMMEDIATE')
        try:
            cols=[r[1] for r in raw.execute('PRAGMA table_info(liquidations)')]
            raw.execute('CREATE TABLE liquidations_new(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),period TEXT NOT NULL,cents INTEGER NOT NULL,quarters INTEGER NOT NULL,created_at TEXT NOT NULL,created_by INTEGER REFERENCES users(id),cutoff_date TEXT)')
            raw.execute('INSERT INTO liquidations_new(id,user_id,period,cents,quarters,created_at,created_by,cutoff_date) SELECT id,user_id,period,cents,quarters,created_at,created_by,'+('cutoff_date' if 'cutoff_date' in cols else 'NULL')+' FROM liquidations')
            raw.execute('DROP TABLE liquidations')
            raw.execute('ALTER TABLE liquidations_new RENAME TO liquidations')
            raw.commit()
        except Exception:
            raw.rollback();raise
    raw.close()
    with conn() as c:
        c.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, name TEXT NOT NULL, pw TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','supervisor','trabajador')), active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY, code TEXT UNIQUE NOT NULL, name TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, supervisor_id INTEGER REFERENCES users(id));
        CREATE TABLE IF NOT EXISTS activities(id INTEGER PRIMARY KEY, project_id INTEGER NOT NULL REFERENCES projects(id), code TEXT NOT NULL, name TEXT NOT NULL, budget_quarters INTEGER NOT NULL DEFAULT 0 CHECK(budget_quarters>=0), active INTEGER NOT NULL DEFAULT 1, UNIQUE(project_id,code));
        CREATE TABLE IF NOT EXISTS assignments(user_id INTEGER REFERENCES users(id), project_id INTEGER REFERENCES projects(id), PRIMARY KEY(user_id,project_id));
        CREATE TABLE IF NOT EXISTS rates(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id), effective_date TEXT NOT NULL, rate_cents INTEGER NOT NULL CHECK(rate_cents>=0), UNIQUE(user_id,effective_date));
        CREATE TABLE IF NOT EXISTS entries(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), activity_id INTEGER NOT NULL REFERENCES activities(id), work_date TEXT NOT NULL, quarters INTEGER NOT NULL CHECK(quarters>0 AND quarters<=48), description TEXT NOT NULL, state TEXT NOT NULL CHECK(state IN ('borrador','pendiente','observado','aprobado','rechazado','liquidado')), note TEXT DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL, approval_by INTEGER REFERENCES users(id), approval_at TEXT, liquidation_id INTEGER, submission_key TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS liquidations(id INTEGER PRIMARY KEY, user_id INTEGER REFERENCES users(id), period TEXT NOT NULL, cents INTEGER NOT NULL, quarters INTEGER NOT NULL, created_at TEXT NOT NULL, created_by INTEGER REFERENCES users(id));
        CREATE TABLE IF NOT EXISTS liquidation_lines(id INTEGER PRIMARY KEY, liquidation_id INTEGER REFERENCES liquidations(id), entry_id INTEGER UNIQUE REFERENCES entries(id), quarters INTEGER NOT NULL, rate_cents INTEGER NOT NULL, amount_cents INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, actor_id INTEGER, action TEXT NOT NULL, target TEXT NOT NULL, target_id INTEGER, detail TEXT, at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS activity_assignments(user_id INTEGER NOT NULL REFERENCES users(id), activity_id INTEGER NOT NULL REFERENCES activities(id), planned_quarters INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(user_id,activity_id));
        CREATE TABLE IF NOT EXISTS import_log(id INTEGER PRIMARY KEY,project_id INTEGER,filename TEXT,processed_at TEXT,actor_id INTEGER,notes TEXT);
        ''')
        columns={x[1] for x in c.execute('PRAGMA table_info(liquidations)').fetchall()}
        project_cols={x[1] for x in c.execute('PRAGMA table_info(projects)').fetchall()}
        if 'status' not in project_cols:
            c.execute("ALTER TABLE projects ADD COLUMN status TEXT NOT NULL DEFAULT 'activo'")
            c.execute("UPDATE projects SET status='desactivado' WHERE active=0")
        c.execute("UPDATE projects SET active=CASE WHEN status='activo' THEN 1 ELSE 0 END WHERE active != CASE WHEN status='activo' THEN 1 ELSE 0 END")
        if 'cutoff_date' not in columns:c.execute('ALTER TABLE liquidations ADD COLUMN cutoff_date TEXT')
        if 'progress_pct' not in {x[1] for x in c.execute('PRAGMA table_info(activities)').fetchall()}:c.execute('ALTER TABLE activities ADD COLUMN progress_pct REAL NOT NULL DEFAULT 0')
        if 'must_change_pw' not in {x[1] for x in c.execute('PRAGMA table_info(users)').fetchall()}:c.execute('ALTER TABLE users ADD COLUMN must_change_pw INTEGER NOT NULL DEFAULT 0')
init()

def norm(t):
    return ''.join(ch for ch in unicodedata.normalize('NFKD',str(t).casefold()) if not unicodedata.combining(ch)).strip()
def login_candidate(full):
    words=re.findall(r'[a-z0-9]+',norm(full))
    return words[-1] if words else 'trabajador'
def unique_login(c,full):
    base=login_candidate(full);candidate=base;index=2
    while c.execute('SELECT 1 FROM users WHERE username=?',(candidate,)).fetchone():
        candidate=f'{base}{index}';index+=1
    return candidate
def seed_password(full):
    words=re.findall(r'[a-z0-9]+',norm(full))
    return words[0] if words else secrets.token_urlsafe(8)
def now():return datetime.now().strftime('%Y-%m-%d %H:%M:%S')
def period():return date.today().strftime('%Y-%m')
def audit(c,action,target,tid=None,detail=''):
    c.execute('INSERT INTO audit(actor_id,action,target,target_id,detail,at) VALUES(?,?,?,?,?,?)',(session.get('uid'),action,target,tid,str(detail)[:500],now()))
def user():
    if not session.get('uid'):return None
    with conn() as c:return c.execute('SELECT * FROM users WHERE id=? AND active=1',(session['uid'],)).fetchone()
def login_required(f):
    @wraps(f)
    def w(*args,**kw):
        if not user():session.clear();return redirect(url_for('login'))
        if user()['must_change_pw'] and request.endpoint not in ('change_password','logout'):return redirect(url_for('change_password'))
        return f(*args,**kw)
    return w
def roles(*allowed):
    def decor(f):
        @wraps(f)
        @login_required
        def w(*a,**kw):
            if user()['role'] not in allowed:abort(403)
            return f(*a,**kw)
        return w
    return decor
def csrf():
    if '_csrf' not in session:session['_csrf']=secrets.token_urlsafe(24)
    return session['_csrf']
@app.before_request
def protect():
    if request.method=='POST' and request.form.get('_csrf') != csrf():abort(400,'Token de seguridad invalido. Recarga la pagina.')

def fmt(cents):return f'S/ {cents/100:,.2f}'
app.jinja_env.globals.update(fmt=fmt,csrf=csrf)
BASE='''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ALTIVIA HH CONTROL</title><style>
:root{font-family:Inter,Segoe UI,system-ui,sans-serif;color:#142f43;background:#f1f5f9}*{box-sizing:border-box}body{margin:0}header{background:#102f43;color:#fff;padding:16px max(24px,calc((100vw - 1450px)/2));display:flex;justify-content:space-between;align-items:center;gap:20px;flex-wrap:wrap;border-bottom:3px solid #00a69a}header strong{font-size:20px;letter-spacing:.03em}header a{color:#d9eff2;text-decoration:none;display:inline-block;margin:3px 8px;font-size:13px}header a:hover{color:#fff;text-decoration:underline}main{max-width:1450px;margin:27px auto;padding:0 24px}.card{background:white;border:1px solid #e3ebef;border-radius:15px;box-shadow:0 3px 14px #18384b0c;padding:22px;margin-bottom:19px}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(205px,1fr));gap:15px}.stat{font-size:29px;font-weight:800;color:#0f857f}.muted{color:#5c7380}label{display:block;font-size:13px;font-weight:700;margin-top:7px}input,select,textarea{display:block;width:100%;border:1px solid #cbd8e0;border-radius:8px;padding:10px;font:inherit;margin:4px 0 8px;background:white;min-width:0}input:focus,select:focus,textarea:focus{outline:2px solid #9bdcd7;border-color:#128880}button,.btn{display:inline-block;border:0;background:#087f78;color:white;padding:10px 16px;border-radius:8px;text-decoration:none;cursor:pointer;font:inherit;margin:3px;font-size:13px;font-weight:600}button:hover,.btn:hover{filter:brightness(1.08)}.warning{background:#fff0d7;padding:12px;border-radius:8px}.error{background:#ffe4e4;padding:12px;border-radius:8px}.ok{background:#ddf3ec;padding:12px;border-radius:8px}table{width:100%;border-collapse:collapse}td,th{padding:10px 11px;border-bottom:1px solid #e8eef1;text-align:left;vertical-align:middle}th{background:#ebf2f5;color:#2e5365;font-size:12px;text-transform:uppercase;letter-spacing:.025em}tr:hover td{background:#fafcfd}.scroll{overflow:auto}small{color:#59717d}.inline{display:inline}.inline button{font-size:12px}.danger{background:#a83238}.secondary{background:#526c80}h1{font-size:29px;margin:0 0 12px}h2{font-size:19px;margin:0 0 15px}a{color:#107970}details summary{cursor:pointer}.table-input{margin:0!important;padding:7px!important;font-size:13px!important}.kpi-label{font-size:12px;color:#637987;font-weight:650}.kpi-value{font-weight:800;font-size:27px;color:#174d5a}.progress-track{height:10px;background:#e5edf1;border-radius:10px;overflow:hidden;min-width:100px}.progress-fill{height:100%;background:#168b84;border-radius:10px}.toolbar{display:flex;gap:12px;align-items:end;flex-wrap:wrap}.toolbar>*{flex:1;min-width:170px}@media(max-width:650px){header{padding:14px}main{padding:0 12px}table{font-size:12px}.card{padding:15px}h1{font-size:23px}}

</style></head><body><header><div><strong>ALTIVIA | HH CONTROL</strong></div><nav>{% if current %}<a href="{{url_for('dashboard')}}">Inicio</a><a href="{{url_for('my_entries')}}">Mis HH</a>{% if current.role in ('supervisor','admin') %}<a href="{{url_for('people_hh')}}">HH por profesional</a>{% endif %}{% if current.role in ('supervisor','admin') %}<a href="{{url_for('approvals')}}">Aprobaciones</a><a href="{{url_for('budget')}}">Presupuestos</a>{% endif %}{% if current.role=='admin' %}<a href="{{url_for('admin')}}">Administración</a><a href="{{url_for('manage_activities')}}">Actividades</a><a href="{{url_for('import_matrix')}}">Importar Excel</a><a href="{{url_for('liquidations')}}">Liquidaciones</a><a href="{{url_for('audit_view')}}">Auditoría</a>{% endif %}<a href="{{url_for('logout')}}">Salir ({{current.name}})</a>{% endif %}</nav></header><main>{% with messages=get_flashed_messages(with_categories=true) %}{% for category,msg in messages %}<div class="{{'error' if category=='error' else 'ok'}}">{{msg}}</div>{% endfor %}{% endwith %}{{content|safe}}</main></body></html>'''
def page(template,**context):
    html=render_template_string(template,**context)
    return render_template_string(BASE,content=html,current=user())

def require_positive_hours(value):
    try:q=Decimal(str(value))*4
    except (InvalidOperation,ValueError):raise ValueError('HH no validas')
    if q!=q.to_integral_value() or not (1<=int(q)<=48):raise ValueError('HH deben estar entre 0.25 y 12, en pasos de 0.25')
    return int(q)

def can_review(c,e,u):
    if e['user_id']==u['id']:return False
    if u['role']=='admin':return True
    if u['role']!='supervisor':return False
    p=c.execute('SELECT p.supervisor_id FROM projects p JOIN activities a ON a.project_id=p.id WHERE a.id=?',(e['activity_id'],)).fetchone()
    return p and p['supervisor_id']==u['id']

@app.route('/setup',methods=['GET','POST'])
def setup():
    with conn() as c:
        if c.execute('SELECT COUNT(*) FROM users').fetchone()[0]:return redirect(url_for('login'))
        if request.method=='POST':
            name=request.form['name'].strip(); username=request.form['username'].strip().lower();pwd=request.form['password']
            if len(pwd)<10 or not name or not username:flash('Nombre, usuario y contraseña (mínimo 10 caracteres) obligatorios.','error')
            else:
                c.execute('INSERT INTO users(username,name,pw,role) VALUES(?,?,?,?)',(username,name,generate_password_hash(pwd),'admin'))
                c.commit();flash('Administrador creado. Ya puedes iniciar sesión.','ok');return redirect(url_for('login'))
    return page('''<div class="card"><h1>Primer inicio - crear administrador</h1><p>Los datos solo se guardarán en esta PC.</p><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><label>Nombre completo</label><input name="name" required><label>Usuario</label><input name="username" required><label>Contraseña (mínimo 10 caracteres)</label><input type="password" name="password" minlength="10" required><button>Crear administrador</button></form></div>''')
@app.route('/login',methods=['GET','POST'])
def login():
    with conn() as c:
        if not c.execute('SELECT COUNT(*) FROM users').fetchone()[0]:return redirect(url_for('setup'))
        if request.method=='POST':
            u=c.execute('SELECT * FROM users WHERE username=? AND active=1',(request.form.get('username','').lower().strip(),)).fetchone()
            if u and check_password_hash(u['pw'],request.form.get('password','')):
                session.clear();session['uid']=u['id'];csrf();return redirect(url_for('dashboard'))
            flash('Credenciales incorrectas.','error')
    return page('''<div class="card" style="max-width:420px;margin:auto"><h1>Ingresar a ALTIVIA</h1><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><label>Usuario</label><input name="username" required><label>Contraseña</label><input name="password" type="password" required><button>Ingresar</button></form></div>''')
@app.route('/change-password',methods=['GET','POST'])
@login_required
def change_password():
    if request.method=='POST':
        pw=request.form.get('password','')
        if len(pw)<10:flash('Elige una contraseña nueva de al menos 10 caracteres.','error')
        elif pw!=request.form.get('repeat'):flash('Las contraseñas no coinciden.','error')
        else:
            with conn() as c:
                c.execute('UPDATE users SET pw=?,must_change_pw=0 WHERE id=?',(generate_password_hash(pw),user()['id']))
                audit(c,'cambio_clave','user',user()['id']);c.commit()
            flash('Contraseña actualizada.','ok');return redirect(url_for('dashboard'))
    return page('''<div class="card" style="max-width:480px;margin:auto"><h1>Actualizar contraseña inicial</h1><p>La contraseña importada del Excel es temporal y debe cambiarse antes de usar la aplicación.</p><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><label>Nueva contraseña (mín. 10 caracteres)</label><input type="password" name="password" minlength="10" required><label>Repetir contraseña</label><input type="password" name="repeat" minlength="10" required><button>Guardar nueva contraseña</button></form></div>''')
@app.route('/logout')
def logout():session.clear();return redirect(url_for('login'))
@app.route('/')
@login_required
def dashboard():
    actor=user()
    start_date=request.args.get('start_date','').strip()
    end_date=request.args.get('end_date','').strip()
    for d in (start_date,end_date):
        if d:
            try: date.fromisoformat(d)
            except ValueError: abort(400,'Fecha inválida')
    if start_date and end_date and start_date>end_date: abort(400,'Fecha inicial posterior a la final')
    with conn() as c:
        if actor['role']=='admin': project_filter='1=1';scope=[]
        elif actor['role']=='supervisor':project_filter='p.supervisor_id=?';scope=[actor['id']]
        else:project_filter='p.id IN (SELECT project_id FROM assignments WHERE user_id=?)';scope=[actor['id']]
        projects=c.execute('SELECT p.id,p.code,p.name,p.status FROM projects p WHERE '+project_filter+' ORDER BY p.code',scope).fetchall()
        allowed={str(x['id']) for x in projects}
        project=request.args.get('project','')
        if project and project not in allowed:abort(403)
        filter_sql=project_filter+(' AND p.id=?' if project else '')
        params=scope+([int(project)] if project else [])
        date_cond=(' AND e.work_date>=?' if start_date else '')+(' AND e.work_date<=?' if end_date else '')
        date_args=([start_date] if start_date else [])+([end_date] if end_date else [])
        budget_rows=c.execute('''SELECT p.id,p.code,p.name,p.status,COALESCE(SUM(a.budget_quarters),0) budget
          FROM projects p LEFT JOIN activities a ON a.project_id=p.id WHERE '''+filter_sql+' GROUP BY p.id ORDER BY p.code',params).fetchall()
        summary=[]
        for p in budget_rows:
            values=c.execute('''SELECT e.state,COALESCE(SUM(e.quarters),0) AS quarters
                FROM entries e JOIN activities a ON a.id=e.activity_id
                WHERE a.project_id=? AND e.state IN ('pendiente','aprobado','liquidado')'''+date_cond+' GROUP BY e.state',[p['id']]+date_args).fetchall()
            v={r['state']:r['quarters'] for r in values}
            summary.append(dict(p)|{'loaded':sum(v.values()),'approved':v.get('aprobado',0)+v.get('liquidado',0),'pending':v.get('pendiente',0),'unpaid':v.get('aprobado',0),'paid':v.get('liquidado',0)})
        if actor['role']=='trabajador':
            entry_scope='e.user_id=?';entry_args=[actor['id']]
        elif actor['role']=='supervisor':
            entry_scope='p.supervisor_id=?';entry_args=[actor['id']]
        else: entry_scope='1=1';entry_args=[]
        if project:entry_scope+=' AND p.id=?';entry_args.append(int(project))
        rows=c.execute('''SELECT e.state,COALESCE(SUM(e.quarters),0) quarters
          FROM entries e JOIN activities a ON a.id=e.activity_id JOIN projects p ON p.id=a.project_id
          WHERE '''+entry_scope+" AND e.state IN ('pendiente','aprobado','liquidado')"+date_cond+' GROUP BY e.state',entry_args+date_args).fetchall()
    counts={r['state']:r['quarters']/4 for r in rows}
    budget=sum(r['budget'] for r in summary)/4
    loaded=sum(r['loaded'] for r in summary)/4
    return page('''<h1>Centro gerencial</h1><p class="muted">Filtra por proyecto, fechas, ambos o consulta el histórico completo. El presupuesto es total y las HH consumidas corresponden al filtro.</p>
    <div class="card"><form method="get" class="toolbar"><div><label>Proyecto (opcional)</label><select name="project"><option value="">Todos los proyectos</option>{% for p in projects %}<option value="{{p.id}}" {{'selected' if selected==p.id|string else ''}}>{{p.code}} · {{p.name}} ({{p.status}})</option>{% endfor %}</select></div><div><label>Desde (opcional)</label><input type="date" name="start_date" value="{{start_date}}"></div><div><label>Hasta (opcional)</label><input type="date" name="end_date" value="{{end_date}}"></div><div><button>Aplicar filtros</button><a class="btn secondary" href="{{url_for('dashboard')}}">Limpiar</a></div></form></div>
    <div class="grid"><div class="card"><div class="kpi-label">Presupuesto total de proyectos seleccionados</div><div class="kpi-value">{{'%.2f'|format(budget)}}</div></div><div class="card"><div class="kpi-label">HH cargadas · filtro</div><div class="kpi-value">{{'%.2f'|format(loaded)}}</div></div><div class="card"><div class="kpi-label">Presupuesto menos HH del filtro</div><div class="kpi-value">{{'%.2f'|format(budget-loaded)}}</div></div><div class="card"><div class="kpi-label">HH por aprobar · filtro</div><div class="kpi-value">{{'%.2f'|format(counts.get('pendiente',0))}}</div></div><div class="card"><div class="kpi-label">HH aprobadas sin liquidar · filtro</div><div class="kpi-value">{{'%.2f'|format(counts.get('aprobado',0))}}</div></div><div class="card"><div class="kpi-label">HH liquidadas · filtro</div><div class="kpi-value">{{'%.2f'|format(counts.get('liquidado',0))}}</div></div></div>
    <div class="card scroll"><h2>Presupuesto vs. HH por proyecto</h2><p class="muted">Con filtro de fechas, las HH y el saldo mostrado se refieren únicamente al periodo seleccionado, no al saldo histórico real del proyecto.</p><table><thead><tr><th>Proyecto</th><th>Estado</th><th>Presupuesto total HH</th><th>HH cargadas (filtro)</th><th>HH aprobadas</th><th>HH por aprobar</th><th>Diferencia HH (filtro)</th><th>Consumo / presupuesto</th></tr></thead><tbody>{% for p in projs %}{% set ratio=100*p.loaded/p.budget if p.budget else 0 %}<tr><td><strong>{{p.code}}</strong><br><small>{{p.name}}</small></td><td>{{p.status|capitalize}}</td><td>{{p.budget/4}}</td><td><strong>{{p.loaded/4}}</strong></td><td>{{p.approved/4}}</td><td>{{p.pending/4}}</td><td><strong>{{(p.budget-p.loaded)/4}}</strong></td><td><div class="progress-track"><div class="progress-fill" style="width:{{[ratio,100]|min}}%;background:{{'#b93737' if ratio>100 else '#168b84'}}"></div></div><small>{{'%.1f'|format(ratio)}}% de consumo en el filtro</small></td></tr>{% endfor %}</tbody></table></div>
    <div class="card"><h2>Accesos rápidos</h2><a class="btn" href="{{url_for('my_entries')}}">Cargar HH</a>{% if role!='trabajador' %}<a class="btn secondary" href="{{url_for('people_hh')}}">HH por profesional</a><a class="btn secondary" href="{{url_for('approvals')}}">Aprobaciones</a><a class="btn secondary" href="{{url_for('budget')}}">Presupuestos EDT</a>{% endif %}{% if role=='admin' %}<a class="btn" href="{{url_for('manage_activities')}}">Gestionar proyectos</a><a class="btn secondary" href="{{url_for('import_matrix')}}">Importar Excel</a><a class="btn secondary" href="{{url_for('liquidations')}}">Liquidaciones</a>{% endif %}</div>''',projects=projects,selected=project,start_date=start_date,end_date=end_date,projs=summary,counts=counts,budget=budget,loaded=loaded,role=actor['role'])

@app.route('/hh-personas')
@roles('admin','supervisor')
def people_hh():
    actor=user(); month=request.args.get('month','')
    if month and (len(month)!=7 or month[4]!='-' or not month[:4].isdigit() or not month[5:].isdigit() or not 1<=int(month[5:])<=12):abort(400)
    with conn() as c:
        projects=c.execute('SELECT id,code,name FROM projects '+('WHERE supervisor_id=?' if actor['role']=='supervisor' else '')+' ORDER BY code',([actor['id']] if actor['role']=='supervisor' else [])).fetchall()
        pid=request.args.get('project','')
        permitted={str(r['id']) for r in projects}
        if pid and pid not in permitted:abort(403)
        conditions=["e.state IN ('pendiente','aprobado','liquidado')"]
        args=[]
        if actor['role']=='supervisor': conditions.append('p.supervisor_id=?');args.append(actor['id'])
        if pid:conditions.append('p.id=?');args.append(int(pid))
        if month:conditions.append('substr(e.work_date,1,7)=?');args.append(month)
        base=' FROM entries e JOIN users u ON u.id=e.user_id JOIN activities a ON a.id=e.activity_id JOIN projects p ON p.id=a.project_id WHERE '+' AND '.join(conditions)
        grouped=c.execute("""SELECT p.id pid,p.code pc,p.name pn,u.id uid,u.name un,SUM(e.quarters) total,SUM(CASE WHEN e.state='liquidado' THEN e.quarters ELSE 0 END) paid,SUM(CASE WHEN e.state='aprobado' THEN e.quarters ELSE 0 END) approved,SUM(CASE WHEN e.state='pendiente' THEN e.quarters ELSE 0 END) pending """+base+' GROUP BY p.id,u.id ORDER BY p.code,u.name',args).fetchall()
        details=c.execute('SELECT p.id pid,u.id uid,a.code ac,a.name an,SUM(e.quarters) total '+base+' GROUP BY p.id,u.id,a.id ORDER BY a.code',args).fetchall()
    lookup={}
    for r in details:lookup.setdefault((r['pid'],r['uid']),[]).append(dict(r))
    total=sum(r['total'] for r in grouped)/4
    return page('''<h1>HH consolidadas por proyecto y profesional</h1><p class="muted">Consulta gerencial. Los totales agregan todas las actividades y pueden desplegarse para ver su distribución.</p><div class="card"><form method="get" class="grid"><div><label>Proyecto</label><select name="project"><option value="">Todos los proyectos</option>{% for p in projects %}<option value="{{p.id}}" {{'selected' if pid==p.id|string else ''}}>{{p.code}} · {{p.name}}</option>{% endfor %}</select></div><div><label>Mes (vacío = histórico)</label><input type="month" name="month" value="{{month}}"></div><div><label>Consultar</label><button>Filtrar</button></div></form></div><div class="grid"><div class="card"><div class="muted">Total de HH registradas en estados de control</div><div class="stat">{{total}} HH</div></div><div class="card"><div class="muted">Profesionales por proyecto</div><div class="stat">{{rows|length}}</div></div></div><div class="card scroll"><h2>Distribución de HH por profesional</h2><table><tr><th>Proyecto</th><th>Profesional</th><th>Total HH</th><th>Por aprobar</th><th>Aprobadas sin liquidar</th><th>Liquidadas</th><th>Detalle</th></tr>{% for r in rows %}<tr><td>{{r.pc}} · {{r.pn}}</td><td><strong>{{r.un}}</strong></td><td><strong>{{r.total/4}}</strong></td><td>{{r.pending/4}}</td><td>{{r.approved/4}}</td><td>{{r.paid/4}}</td><td><details><summary>Actividades</summary><table><tr><th>EDT</th><th>Actividad</th><th>HH</th></tr>{% for a in detail.get((r.pid,r.uid),[]) %}<tr><td>{{a.ac}}</td><td>{{a.an}}</td><td>{{a.total/4}}</td></tr>{% endfor %}</table></details></td></tr>{% endfor %}</table>{% if not rows %}<p>No hay HH para esos filtros.</p>{% endif %}</div>''',projects=projects,pid=int(pid) if pid else '',month=month,rows=grouped,detail=lookup,total=total)

@app.route('/my',methods=['GET','POST'])
@login_required
def my_entries():
    u=user()
    with conn() as c:
        if request.method=='POST':
            try:
                a=int(request.form['activity']);q=require_positive_hours(request.form['hours']);d=date.fromisoformat(request.form['work_date']).isoformat(); desc=request.form['description'].strip()
                if len(desc)<5:raise ValueError('Descripción de al menos 5 caracteres')
                allowed=c.execute('''SELECT 1 FROM activities a JOIN projects p ON p.id=a.project_id JOIN assignments x ON x.project_id=p.id WHERE a.id=? AND a.active=1 AND p.active=1 AND p.status='activo' AND x.user_id=?''',(a,u['id'])).fetchone()
                if not allowed:raise ValueError('Actividad no asignada')
                restricted=c.execute('SELECT 1 FROM activity_assignments aa JOIN activities ax ON ax.id=aa.activity_id WHERE ax.project_id=(SELECT project_id FROM activities WHERE id=?) LIMIT 1',(a,)).fetchone()
                if restricted and not c.execute('SELECT 1 FROM activity_assignments WHERE user_id=? AND activity_id=?',(u['id'],a)).fetchone():raise ValueError('No estás autorizado en esta actividad según la matriz de responsabilidades')
                if d>date.today().isoformat():raise ValueError('No se permiten fechas futuras')
                if (date.today()-date.fromisoformat(d)).days>45:raise ValueError('Fecha demasiado antigua; consulta con el administrador')
                daily=c.execute("SELECT COALESCE(SUM(quarters),0) FROM entries WHERE user_id=? AND work_date=? AND state!='rechazado'",(u['id'],d)).fetchone()[0]
                if daily+q>48:raise ValueError('Excede el máximo configurable de 12 HH por día')
                key=request.form.get('submission_key','')
                if not key:raise ValueError('Identificador de formulario inválido')
                c.execute('INSERT OR IGNORE INTO entries(user_id,activity_id,work_date,quarters,description,state,created_at,updated_at,submission_key) VALUES(?,?,?,?,?,?,?,?,?)',(u['id'],a,d,q,desc,'pendiente',now(),now(),key))
                audit(c,'registro','entry',c.execute('SELECT last_insert_rowid()').fetchone()[0]);c.commit()
                flash('Registro enviado a aprobación.','ok');return redirect(url_for('my_entries'))
            except (ValueError,KeyError,sqlite3.IntegrityError) as ex:flash(str(ex),'error')
        acts=c.execute("""SELECT a.id,p.code||' · '||p.name||' / '||a.code||' · '||a.name label,p.id project_id,a.code ac,a.name an FROM activities a JOIN projects p ON p.id=a.project_id JOIN assignments x ON x.project_id=p.id WHERE x.user_id=? AND a.active=1 AND p.active=1 AND p.status='activo' AND (NOT EXISTS(SELECT 1 FROM activity_assignments z JOIN activities aa ON aa.id=z.activity_id WHERE aa.project_id=p.id) OR EXISTS(SELECT 1 FROM activity_assignments z WHERE z.activity_id=a.id AND z.user_id=?)) ORDER BY p.code,a.code""",(u['id'],u['id'])).fetchall()
        projects=c.execute('SELECT DISTINCT p.id,p.code,p.name FROM projects p JOIN activities a ON a.project_id=p.id JOIN assignments x ON x.project_id=p.id WHERE x.user_id=? AND a.active=1 AND p.active=1 ORDER BY p.code',(u['id'],)).fetchall()
        rows=c.execute('''SELECT e.*,a.code ac,p.code pc,p.name pn FROM entries e JOIN activities a ON a.id=e.activity_id JOIN projects p ON p.id=a.project_id WHERE e.user_id=? ORDER BY e.work_date DESC,e.id DESC LIMIT 200''',(u['id'],)).fetchall()
    return page('''<h1>Mis horas hombre</h1><div class="card"><h2>Nuevo registro</h2>{% if not acts %}<div class="warning">El administrador debe asignarte un proyecto antes de registrar HH.</div>{% else %}<form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="submission_key" value="{{key}}"><div class="grid"><div><label>Fecha de trabajo</label><input name="work_date" type="date" value="{{today}}" required></div><div><label>Proyecto</label><select id="myProject" onchange="filterMyActivities()"><option value="">Seleccionar proyecto</option>{% for p in projects %}<option value="{{p.id}}">{{p.code}} · {{p.name}}</option>{% endfor %}</select><label>Actividad EDT/WBS</label><select id="myActivity" name="activity" required><option value="">Seleccionar actividad</option>{% for a in acts %}<option data-project="{{a.project_id}}" value="{{a.id}}">{{a.ac}} · {{a.an}}</option>{% endfor %}</select></div><div><label>Horas hombre (0.25 a 12)</label><input name="hours" type="number" step="0.25" min="0.25" max="12" required></div></div><label>Descripción técnica</label><textarea name="description" maxlength="1000" required></textarea><button>Enviar para aprobación</button></form><script>function filterMyActivities(){const p=document.getElementById('myProject').value;const a=document.getElementById('myActivity');a.value='';Array.from(a.options).forEach((o,i)=>{if(i)o.hidden=o.disabled=o.dataset.project!==p});}filterMyActivities();</script>{% endif %}</div><div class="card scroll"><h2>Historial (últimos 200)</h2><table><tr><th>Fecha</th><th>Proyecto</th><th>Actividad</th><th>HH</th><th>Estado</th><th>Descripción / Observación</th><th>Acción</th></tr>{% for r in rows %}<tr><td>{{r.work_date}}</td><td>{{r.pc}}</td><td>{{r.ac}}</td><td>{{r.quarters/4}}</td><td>{{r.state}}</td><td>{{r.description}}{% if r.note %}<hr><small>{{r.note}}</small>{% endif %}</td><td>{% if r.state=='observado' %}<a href="{{url_for('correct',eid=r.id)}}">Corregir</a>{% endif %}</td></tr>{% endfor %}</table></div>''',acts=acts,projects=projects,rows=rows,key=secrets.token_hex(16),today=date.today().isoformat())
@app.route('/correct/<int:eid>',methods=['GET','POST'])
@login_required
def correct(eid):
    with conn() as c:
        e=c.execute("SELECT * FROM entries WHERE id=? AND user_id=? AND state='observado'",(eid,user()['id'])).fetchone()
        if not e:abort(404)
        if request.method=='POST':
            try:
                q=require_positive_hours(request.form['hours']);desc=request.form['description'].strip()
                if len(desc)<5:raise ValueError('Descripción insuficiente')
                c.execute("UPDATE entries SET quarters=?,description=?,state='pendiente',note='',updated_at=?,approval_by=NULL,approval_at=NULL WHERE id=?",(q,desc,now(),eid));audit(c,'correccion','entry',eid);c.commit();flash('Corrección reenviada.','ok');return redirect(url_for('my_entries'))
            except ValueError as ex:flash(str(ex),'error')
    return page('''<div class="card"><h1>Corregir HH observadas</h1><p>Observación: {{e.note}}</p><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><label>HH</label><input type="number" name="hours" min="0.25" max="12" step="0.25" value="{{e.quarters/4}}" required><label>Descripción</label><textarea name="description" required>{{e.description}}</textarea><button>Reenviar</button></form></div>''',e=e)
@app.route('/approvals',methods=['GET','POST'])
@roles('admin','supervisor')
def approvals():
    u=user()
    with conn() as c:
        if request.method=='POST':
            eid=int(request.form.get('entry_id',0));decision=request.form.get('decision');note=request.form.get('note','').strip()
            e=c.execute("SELECT * FROM entries WHERE id=? AND state='pendiente'",(eid,)).fetchone()
            if not e or not can_review(c,e,u):abort(403)
            if decision not in ('aprobado','observado','rechazado'):abort(400)
            if decision!='aprobado' and len(note)<4:flash('Observación obligatoria (4 caracteres mínimo).','error')
            else:
                c.execute('UPDATE entries SET state=?,note=?,approval_by=?,approval_at=?,updated_at=? WHERE id=?',(decision,note,u['id'],now(),now(),eid));audit(c,'revision_'+decision,'entry',eid,note);c.commit();flash('Revisión guardada.','ok')
            return redirect(url_for('approvals'))
        rows=c.execute('''SELECT e.*,u.name person,p.code pc,a.code ac,a.name an FROM entries e JOIN users u ON u.id=e.user_id JOIN activities a ON a.id=e.activity_id JOIN projects p ON p.id=a.project_id WHERE e.state='pendiente' AND e.user_id!=? AND (?='admin' OR p.supervisor_id=?) ORDER BY e.work_date,e.id''',(u['id'],u['role'],u['id'])).fetchall()
    return page('''<h1>Aprobación de HH</h1><div class="card scroll"><table><tr><th>Trabajador</th><th>Fecha</th><th>Proyecto / actividad</th><th>HH</th><th>Descripción</th><th>Decisión</th></tr>{% for r in rows %}<tr><td>{{r.person}}</td><td>{{r.work_date}}</td><td>{{r.pc}} / {{r.ac}} · {{r.an}}</td><td>{{r.quarters/4}}</td><td>{{r.description}}</td><td><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="entry_id" value="{{r.id}}"><input name="note" placeholder="Motivo si observas/rechazas"><button name="decision" value="aprobado">Aprobar</button><button class="secondary" name="decision" value="observado">Observar</button><button class="danger" name="decision" value="rechazado">Rechazar</button></form></td></tr>{% endfor %}</table>{% if not rows %}<p>No hay HH pendientes para revisar.</p>{% endif %}</div>''',rows=rows)
@app.route('/budget')
@roles('admin','supervisor')
def budget():
    u=user()
    with conn() as c:
        rows=c.execute('''SELECT p.code pc,p.name pn,a.code ac,a.name an,a.budget_quarters budget,COALESCE(SUM(CASE WHEN e.state IN ('aprobado','liquidado') THEN e.quarters ELSE 0 END),0) approved,COALESCE(SUM(CASE WHEN e.state='pendiente' THEN e.quarters ELSE 0 END),0) pending FROM activities a JOIN projects p ON p.id=a.project_id LEFT JOIN entries e ON e.activity_id=a.id WHERE (?='admin' OR p.supervisor_id=?) GROUP BY a.id ORDER BY p.code,a.code''',(u['role'],u['id'])).fetchall()
    return page('''<h1>Control presupuestal por actividad</h1><div class="card scroll"><table><tr><th>Proyecto</th><th>Actividad WBS</th><th>Presupuesto</th><th>Aprobadas + liquidadas</th><th>Pendientes</th><th>Saldo validado</th><th>Uso total (incl. pendiente)</th></tr>{% for r in rows %}{% set used=r.approved+r.pending %}<tr><td>{{r.pc}} · {{r.pn}}</td><td>{{r.ac}} · {{r.an}}</td><td>{{r.budget/4}}</td><td>{{r.approved/4}}</td><td>{{r.pending/4}}</td><td>{{(r.budget-r.approved)/4}}</td><td><strong style="color:{{'#b52b31' if used>r.budget else '#af6a08' if r.budget>0 and used*10>=r.budget*8 else '#137468'}}">{{'%.1f'|format(used*100/r.budget) if r.budget else 'Sin presupuesto'}}{% if r.budget %}%{% endif %}</strong></td></tr>{% endfor %}</table></div>''',rows=rows)
@app.route('/admin',methods=['GET','POST'])
@roles('admin')
def admin():
    with conn() as c:
        if request.method=='POST':
            kind=request.form.get('kind')
            try:
                if kind=='user':
                    name=request.form['name'].strip();role=request.form.get('role','trabajador')
                    if not name or role not in ('admin','supervisor','trabajador'):raise ValueError('Nombre o rol inválido')
                    username=unique_login(c,name);pwd=seed_password(name)
                    c.execute('INSERT INTO users(username,name,pw,role,must_change_pw) VALUES(?,?,?,?,1)',(username,name,generate_password_hash(pwd),role))
                    flash(f'Cuenta creada: usuario {username}. Contraseña inicial temporal: primer nombre. Debe cambiarla al ingresar.','ok')
                elif kind=='user_edit':
                    uid=int(request.form['user_id']);role=request.form['role'];name=request.form['name'].strip();username=request.form['username'].strip().lower()
                    if role not in ('admin','supervisor','trabajador') or not name or not username:raise ValueError('Datos de usuario inválidos')
                    if uid==user()['id'] and role!='admin':raise ValueError('No puedes quitarte tu propio rol administrador')
                    if not c.execute('SELECT 1 FROM users WHERE id=?',(uid,)).fetchone():raise ValueError('Usuario inexistente')
                    c.execute('UPDATE users SET name=?,username=?,role=? WHERE id=?',(name,username,role,uid))
                    audit(c,'editar_usuario','user',uid,f'rol={role}')
                elif kind=='user_disable':
                    uid=int(request.form['user_id']);current=c.execute('SELECT active FROM users WHERE id=?',(uid,)).fetchone()
                    if not current:raise ValueError('Usuario inexistente')
                    if uid==user()['id']:raise ValueError('No puedes desactivar tu propia cuenta')
                    c.execute('UPDATE users SET active=? WHERE id=?',(0 if current['active'] else 1,uid))
                    audit(c,'cambiar_estado_usuario','user',uid)
                elif kind=='user_password':
                    uid=int(request.form['user_id']);pw=request.form.get('password','')
                    if len(pw)<10:raise ValueError('La contraseña de restablecimiento debe tener al menos 10 caracteres')
                    c.execute('UPDATE users SET pw=?,must_change_pw=1 WHERE id=?',(generate_password_hash(pw),uid))
                    audit(c,'reinicio_clave','user',uid)
                elif kind=='activity_edit':
                    aid=int(request.form['activity_id']);q=require_budget(request.form['budget'])
                    if not request.form['code'].strip() or not request.form['name'].strip():raise ValueError('Actividad y código obligatorios')
                    c.execute('UPDATE activities SET code=?,name=?,budget_quarters=? WHERE id=?',(request.form['code'].strip().upper(),request.form['name'].strip(),q,aid))
                    audit(c,'editar_actividad','activity',aid)
                elif kind=='activity_disable':
                    aid=int(request.form['activity_id']);row=c.execute('SELECT active FROM activities WHERE id=?',(aid,)).fetchone()
                    if not row:raise ValueError('Actividad inexistente')
                    c.execute('UPDATE activities SET active=? WHERE id=?',(0 if row['active'] else 1,aid))
                    audit(c,'cambiar_estado_actividad','activity',aid)
                elif kind=='project_status':
                    pid=int(request.form['project_id'])
                    status=request.form.get('status','')
                    if status not in ('activo','desactivado','terminado'):raise ValueError('Estado de proyecto inválido')
                    old=c.execute('SELECT status FROM projects WHERE id=?',(pid,)).fetchone()
                    if not old:raise ValueError('Proyecto inexistente')
                    c.execute('UPDATE projects SET status=?,active=? WHERE id=?',(status,1 if status=='activo' else 0,pid))
                    audit(c,'cambio_estado_proyecto','project',pid,f"{old['status']} -> {status}")
                elif kind=='project':
                    sid=int(request.form['supervisor_id']) if request.form['supervisor_id'] else None
                    if sid and not c.execute("SELECT 1 FROM users WHERE id=? AND role IN ('admin','supervisor')",(sid,)).fetchone():raise ValueError('Supervisor inválido')
                    c.execute('INSERT INTO projects(code,name,supervisor_id) VALUES(?,?,?)',(request.form['code'].strip().upper(),request.form['name'].strip(),sid))
                elif kind=='activity':
                    q=require_budget(request.form['budget'])
                    pid=int(request.form['project_id'])
                    c.execute('INSERT INTO activities(project_id,code,name,budget_quarters) VALUES(?,?,?,?)',(pid,request.form['code'].strip().upper(),request.form['name'].strip(),q))
                elif kind=='assignment':c.execute('INSERT INTO assignments(user_id,project_id) VALUES(?,?)',(int(request.form['user_id']),int(request.form['project_id'])))
                elif kind in ('rate_edit','rate_delete'):
                    rid=int(request.form['rate_id']);old=c.execute('SELECT * FROM rates WHERE id=?',(rid,)).fetchone()
                    if not old:raise ValueError('Tarifa no encontrada')
                    if kind=='rate_delete':
                        c.execute('DELETE FROM rates WHERE user_id=?',(old['user_id'],))
                    else:
                        cents=int((Decimal(request.form['rate'])*100).quantize(Decimal('1')))
                        if cents<0:raise ValueError('Tarifa negativa')
                        c.execute('UPDATE rates SET rate_cents=? WHERE id=?',(cents,rid))
                        c.execute('DELETE FROM rates WHERE user_id=? AND id<>?',(old['user_id'],rid))
                    audit(c,kind,'rate',rid,f"anterior={dict(old)}; los importes ya liquidados permanecen congelados")
                elif kind=='rate':
                    uid=int(request.form['user_id'])
                    cents=int((Decimal(request.form['rate'])*100).quantize(Decimal('1')))
                    if cents<0:raise ValueError('Tarifa inválida')
                    existing=c.execute('SELECT id FROM rates WHERE user_id=? ORDER BY effective_date DESC,id DESC LIMIT 1',(uid,)).fetchone()
                    if existing:
                        c.execute('UPDATE rates SET rate_cents=? WHERE id=?',(cents,existing['id']))
                        c.execute('DELETE FROM rates WHERE user_id=? AND id<>?',(uid,existing['id']))
                    else:
                        c.execute('INSERT INTO rates(user_id,effective_date,rate_cents) VALUES(?,?,?)',(uid,'1900-01-01',cents))
                else:raise ValueError('Operación desconocida')
                audit(c,'alta_'+kind,kind);c.commit();flash('Información guardada.','ok');return redirect(url_for('manage_activities',project_id=request.form.get('project_id') or request.form.get('return_project'))) if request.form.get('return_to')=='activities' else redirect(url_for('admin'))
            except (ValueError,KeyError,sqlite3.IntegrityError,InvalidOperation) as ex:flash(f'No se pudo guardar: {ex}','error')
        users=c.execute('SELECT * FROM users ORDER BY name').fetchall()
        projs=c.execute('SELECT * FROM projects ORDER BY code').fetchall()
        acts=c.execute('SELECT a.*,p.code pc FROM activities a JOIN projects p ON p.id=a.project_id ORDER BY p.code,a.code').fetchall()
        assign=c.execute('SELECT u.name un,p.code pc FROM assignments x JOIN users u ON u.id=x.user_id JOIN projects p ON p.id=x.project_id ORDER BY u.name').fetchall()
        rates=c.execute('SELECT r.*,u.name un FROM rates r JOIN users u ON u.id=r.user_id WHERE r.id=(SELECT rr.id FROM rates rr WHERE rr.user_id=r.user_id ORDER BY rr.effective_date DESC,rr.id DESC LIMIT 1) ORDER BY u.name').fetchall()
    return page('''<h1>Administración de ALTIVIA</h1><p class="muted">Configura usuarios, proyectos y actividades antes de empezar. Las tarifas son estándares por persona y pueden editarse. Las liquidaciones cerradas mantienen sus importes originales.</p><div class="grid"><div class="card"><h2>1. Crear usuario</h2><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="user"><label>Nombre</label><input name="name" required><p class="muted">El usuario se genera del apellido. La clave inicial temporal es el primer nombre y se cambiará al ingresar.</p><label>Rol</label><select name="role"><option value="trabajador">Trabajador</option><option value="supervisor">Supervisor</option><option value="admin">Administrador</option></select><button>Crear usuario</button></form></div><div class="card"><h2>2. Crear proyecto</h2><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="project"><label>Código único</label><input name="code" placeholder="ALT-001" required><label>Nombre</label><input name="name" required><label>Supervisor</label><select name="supervisor_id"><option value="">Sin supervisor</option>{% for u in users if u.role in ('admin','supervisor') %}<option value="{{u.id}}">{{u.name}}</option>{% endfor %}</select><button>Crear proyecto</button></form></div><div class="card"><h2>3. Crear actividad WBS</h2><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="activity"><label>Proyecto</label><select name="project_id">{% for p in projs %}<option value="{{p.id}}">{{p.code}} · {{p.name}}</option>{% endfor %}</select><label>Código EDT/WBS</label><input name="code" required><label>Actividad</label><input name="name" required><label>Presupuesto HH</label><input type="number" name="budget" min="0" step="0.25" required><button>Crear actividad</button></form></div><div class="card"><h2>4. Asignar trabajador</h2><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="assignment"><label>Trabajador</label><select name="user_id">{% for u in users %}<option value="{{u.id}}">{{u.name}}</option>{% endfor %}</select><label>Proyecto</label><select name="project_id">{% for p in projs %}<option value="{{p.id}}">{{p.code}} · {{p.name}}</option>{% endfor %}</select><button>Asignar</button></form></div><div class="card"><h2>5. Tarifa estándar por trabajador</h2><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="rate"><label>Trabajador</label><select name="user_id">{% for u in users %}<option value="{{u.id}}">{{u.name}}</option>{% endfor %}</select><label>Soles por HH</label><input type="number" name="rate" step="0.01" min="0" required><button>Guardar tarifa</button></form></div></div><div class="card scroll"><h2>Equipo de trabajo · tabla editable</h2><p class="muted">Una fila por usuario. Los cambios se guardan por fila; desactivar conserva su historial.</p><table style="min-width:850px"><tr><th>Nombre completo</th><th>Usuario</th><th>Rol</th><th>Estado</th><th>Guardar</th><th>Acciones</th></tr>{% for u in users %}<tr><td><form id="user-{{u.id}}" method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="user_edit"><input type="hidden" name="user_id" value="{{u.id}}"></form><input form="user-{{u.id}}" class="table-input" name="name" value="{{u.name}}" required></td><td><input form="user-{{u.id}}" class="table-input" name="username" value="{{u.username}}" required></td><td><select form="user-{{u.id}}" class="table-input" name="role">{% for r in ['trabajador','supervisor','admin'] %}<option value="{{r}}" {{'selected' if u.role==r else ''}}>{{r|capitalize}}</option>{% endfor %}</select></td><td>{{'Activo' if u.active else 'Inactivo'}}</td><td><button form="user-{{u.id}}">Guardar</button></td><td><form method="post" class="inline" onsubmit="return confirm('¿Cambiar acceso?')"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="user_disable"><input type="hidden" name="user_id" value="{{u.id}}"><button class="{{'danger' if u.active else 'secondary'}}">{{'Desactivar' if u.active else 'Reactivar'}}</button></form><details><summary>Contraseña</summary><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="user_password"><input type="hidden" name="user_id" value="{{u.id}}"><input type="password" name="password" minlength="10" placeholder="Nueva contraseña" required><button>Restablecer</button></form></details></td></tr>{% endfor %}</table></div><div class="card"><h2>Proyectos y actividades</h2><p class="muted">Accede por proyecto a una tabla compacta editable: búsqueda, HH presupuestadas y edición en filas.</p><a class="btn" href="{{url_for('manage_activities')}}">Abrir gestor de proyectos y actividades →</a></div><div class="card scroll"><h2>Asignaciones</h2><table>{% for x in assign %}<tr><td>{{x.un}}</td><td>{{x.pc}}</td></tr>{% endfor %}</table></div><div class="card scroll"><h2>Tarifas estándar (solo administración)</h2><table><tr><th>Trabajador</th><th>Tarifa y acciones</th></tr>{% for r in rates %}<tr><td>{{r.un}}</td><td><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="rate_edit"><input type="hidden" name="rate_id" value="{{r.id}}"><input type="number" step="0.01" min="0" name="rate" value="{{'%.2f'|format(r.rate_cents/100)}}" required><button>Guardar</button></form><form method="post" onsubmit="return confirm('¿Quitar tarifa del trabajador? Las liquidaciones cerradas no cambiarán.')"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="kind" value="rate_delete"><input type="hidden" name="rate_id" value="{{r.id}}"><button class="danger">Eliminar</button></form></td></tr>{% endfor %}</table></div>''',users=users,projs=projs,acts=acts,assign=assign,rates=rates,today=date.today().isoformat())

@app.route('/admin/actividades')
@roles('admin')
def manage_activities():
    chosen=request.args.get('project_id',type=int)
    with conn() as c:
        projects=c.execute('SELECT p.*,COUNT(a.id) n, COALESCE(SUM(a.budget_quarters),0) total FROM projects p LEFT JOIN activities a ON a.project_id=p.id GROUP BY p.id ORDER BY p.code').fetchall()
        if not chosen and projects:chosen=projects[0]['id']
        if chosen and not any(p['id']==chosen for p in projects):abort(404)
        rows=c.execute("SELECT a.*, COALESCE(SUM(CASE WHEN e.state IN ('pendiente','aprobado','liquidado') THEN e.quarters ELSE 0 END),0) used FROM activities a LEFT JOIN entries e ON e.activity_id=a.id WHERE a.project_id=? GROUP BY a.id ORDER BY a.code",(chosen,)).fetchall() if chosen else []
        # Native activity forms remain served by /admin; this view is project-focused.
    return page("""<h1>Gestión de proyectos y actividades</h1><div class='card'><a href='{{url_for("admin")}}'>← Volver a Administración</a><h2 style='margin-top:16px'>Seleccionar proyecto</h2><form method='get'><div class='grid'><div><label>Proyecto</label><select name='project_id' onchange='this.form.submit()'>{% for p in projects %}<option value='{{p.id}}' {{'selected' if p.id==chosen else ''}}>{{p.code}} — {{p.name}} ({{p.n}} actividades)</option>{% endfor %}</select></div><div><label>Buscar actividad</label><input id='searchActivities' type='search' placeholder='Código EDT o descripción...' oninput='filterActivities()'></div></div></form></div>
    {% if chosen %}<div class='grid'><div class='card'><small>Actividades</small><div class='stat'>{{rows|length}}</div></div><div class='card'><small>HH presupuestadas</small><div class='stat'>{{(rows|sum(attribute='budget_quarters'))/4}}</div></div><div class='card'><small>HH registradas / aprobadas</small><div class='stat'>{{(rows|sum(attribute='used'))/4}}</div></div></div>
    <div class='card'><h2>Estado del proyecto</h2><p class='muted'>Al desactivar o terminar un proyecto, los trabajadores no podrán registrar HH nuevas. Se conservan los registros históricos y las liquidaciones.</p><form method='post' action='{{url_for("admin")}}' onsubmit='return confirm("¿Cambiar el estado del proyecto? Esto bloqueará nuevas HH si no está Activo.")'><input type='hidden' name='_csrf' value='{{csrf()}}'><input type='hidden' name='kind' value='project_status'><input type='hidden' name='project_id' value='{{chosen}}'><input type='hidden' name='return_to' value='activities'><label>Estado actual</label><select name='status'><option value='activo' {{'selected' if chosen_status=='activo' else ''}}>Activo</option><option value='desactivado' {{'selected' if chosen_status=='desactivado' else ''}}>Desactivado</option><option value='terminado' {{'selected' if chosen_status=='terminado' else ''}}>Terminado</option></select><button>Guardar estado del proyecto</button></form></div>
    <div class='card'><details><summary><strong>+ Agregar nueva actividad al proyecto</strong></summary><form method='post' action='{{url_for("admin")}}'><input type='hidden' name='_csrf' value='{{csrf()}}'><input type='hidden' name='kind' value='activity'><input type='hidden' name='project_id' value='{{chosen}}'><input type='hidden' name='return_to' value='activities'><div class='grid'><div><label>EDT/WBS</label><input name='code' required></div><div><label>Actividad</label><input name='name' required></div><div><label>Presupuesto HH</label><input type='number' name='budget' step='.25' min='0' required></div></div><button>Agregar actividad</button></form></details></div>
    <div class='card scroll'><h2>Actividades del proyecto</h2><p class='muted'>Edita directamente cada fila y pulsa Guardar. Usa la búsqueda para encontrar tareas. Desactivar conserva registros históricos.</p><table id='activityTable' style='min-width:950px'><thead><tr><th>EDT/WBS</th><th>Actividad / Entregable</th><th>Presupuesto HH</th><th>Consumidas HH</th><th>Estado</th><th>Acciones</th></tr></thead><tbody>{% for a in rows %}<tr class='activity-row' data-key='{{(a.code ~ " " ~ a.name)|lower}}'><td><form id='form-{{a.id}}' method='post' action='{{url_for("admin")}}'><input type='hidden' name='_csrf' value='{{csrf()}}'><input type='hidden' name='kind' value='activity_edit'><input type='hidden' name='activity_id' value='{{a.id}}'><input type='hidden' name='return_project' value='{{chosen}}'><input type='hidden' name='return_to' value='activities'></form><input form='form-{{a.id}}' name='code' value='{{a.code}}' required style='min-width:90px'></td><td><input form='form-{{a.id}}' name='name' value='{{a.name}}' required style='min-width:270px'></td><td><input form='form-{{a.id}}' type='number' name='budget' min='0' step='.25' value='{{a.budget_quarters/4}}' required style='min-width:100px'></td><td><strong>{{a.used/4}}</strong></td><td>{{'Activa' if a.active else 'Inactiva'}}</td><td style='white-space:nowrap'><button form='form-{{a.id}}'>Guardar</button><form method='post' action='{{url_for("admin")}}' style='display:inline' onsubmit='return confirm("¿Cambiar estado de la actividad? Se conservarán los registros anteriores.")'><input type='hidden' name='_csrf' value='{{csrf()}}'><input type='hidden' name='kind' value='activity_disable'><input type='hidden' name='activity_id' value='{{a.id}}'><input type='hidden' name='return_project' value='{{chosen}}'><input type='hidden' name='return_to' value='activities'><button class='{{"danger" if a.active else "secondary"}}'>{{'Desactivar' if a.active else 'Reactivar'}}</button></form></td></tr>{% endfor %}</tbody></table><p id='noMatches' class='muted' style='display:none'>No se encontraron actividades.</p></div>{% else %}<div class='card'>No hay proyectos. Crea un proyecto desde Administración.</div>{% endif %}
    <script>function filterActivities(){const q=document.getElementById('searchActivities').value.toLowerCase().trim();let n=0;document.querySelectorAll('.activity-row').forEach(r=>{let shown=r.dataset.key.includes(q);r.style.display=shown?'':'none';if(shown)n++});document.getElementById('noMatches').style.display=n?'none':'block'}</script>
    """,projects=projects,chosen=chosen,chosen_status=next((p['status'] for p in projects if p['id']==chosen),'activo'),rows=rows)

def require_budget(value):
    q=Decimal(value)*4
    if q!=q.to_integral_value() or q<0:raise ValueError('El presupuesto HH debe ser positivo y múltiplo de 0.25')
    return int(q)

def rate_for(c,uid,work_date):
    r=c.execute('SELECT rate_cents FROM rates WHERE user_id=? ORDER BY effective_date DESC,id DESC LIMIT 1',(uid,)).fetchone()
    return r['rate_cents'] if r else None

def liquidate(c,uid,month,actor,cutoff=None):
    cutoff=cutoff or (date.fromisoformat(month+'-01').replace(day=28)+timedelta(days=4)).replace(day=1)-timedelta(days=1)
    if isinstance(cutoff,date):cutoff=cutoff.isoformat()
    if cutoff[:7]!=month:raise ValueError('La fecha de corte debe pertenecer al periodo seleccionado')
    entries=c.execute("SELECT * FROM entries WHERE user_id=? AND substr(work_date,1,7)=? AND work_date<=? AND date(approval_at)<=? AND state='aprobado' AND liquidation_id IS NULL ORDER BY work_date,id",(uid,month,cutoff,cutoff)).fetchall()
    if not entries:raise ValueError('No existen HH aprobadas sin liquidar hasta esa fecha')
    lines=[]
    for e in entries:
        cents=rate_for(c,uid,e['work_date'])
        if cents is None:raise ValueError('No existe tarifa estándar para este trabajador. Configúrala desde Administración.')
        amount=int((Decimal(e['quarters'])*Decimal(cents)/4).quantize(Decimal('1'),rounding='ROUND_HALF_UP'))
        lines.append((e['id'],e['quarters'],cents,amount))
    total=sum(r[3] for r in lines);quarters=sum(r[1] for r in lines)
    cur=c.execute('INSERT INTO liquidations(user_id,period,cents,quarters,created_at,created_by,cutoff_date) VALUES(?,?,?,?,?,?,?)',(uid,month,total,quarters,now(),actor,cutoff));lid=cur.lastrowid
    for eid,q,r,a in lines:
        c.execute('INSERT INTO liquidation_lines(liquidation_id,entry_id,quarters,rate_cents,amount_cents) VALUES(?,?,?,?,?)',(lid,eid,q,r,a))
        if c.execute("UPDATE entries SET state='liquidado',liquidation_id=?,updated_at=? WHERE id=? AND state='aprobado' AND liquidation_id IS NULL",(lid,now(),eid)).rowcount!=1:raise ValueError('Conflicto de liquidación')
    audit(c,'liquidacion','liquidation',lid,f'{month}, corte {cutoff}, {quarters/4} HH, {total} centimos')
    return lid

@app.route('/liquidations',methods=['GET','POST'])
@roles('admin')
def liquidations():
    month=request.values.get('month',period());cutoff=request.values.get('cutoff',str(date.today()) if month==period() else month+'-15')
    try:
        dt=date.fromisoformat(month+'-01')
        if dt.strftime('%Y-%m')!=month:raise ValueError()
        cutoff=date.fromisoformat(cutoff).isoformat()
        if cutoff[:7]!=month:raise ValueError()
    except ValueError:abort(400,'Periodo inválido')
    with conn() as c:
        if request.method=='POST':
            try:
                uid=int(request.form['uid'])
                if not c.execute('SELECT 1 FROM users WHERE id=?',(uid,)).fetchone():raise ValueError('Trabajador inválido')
                lid=liquidate(c,uid,month,user()['id'],cutoff);c.commit();flash(f'Liquidación #{lid} generada.','ok')
            except (ValueError,sqlite3.IntegrityError) as ex:c.rollback();flash(str(ex),'error')
            return redirect(url_for('liquidations',month=month,cutoff=cutoff))
        available=c.execute("""SELECT u.id,u.name,COUNT(e.id) n,SUM(e.quarters) q FROM users u JOIN entries e ON e.user_id=u.id WHERE substr(e.work_date,1,7)=? AND e.work_date<=? AND date(e.approval_at)<=? AND e.state='aprobado' AND e.liquidation_id IS NULL GROUP BY u.id ORDER BY u.name""",(month,cutoff,cutoff)).fetchall()
        rows=c.execute('SELECT l.*,u.name person FROM liquidations l JOIN users u ON u.id=l.user_id WHERE l.period=? ORDER BY u.name',(month,)).fetchall()
    return page('''<h1>Liquidaciones mensuales</h1><div class="card"><form method="get"><label>Periodo (AAAA-MM)</label><input type="month" name="month" value="{{month}}"><label>Fecha de corte (inclusive)</label><input type="date" name="cutoff" value="{{cutoff}}" required><button>Consultar</button></form><p>Se liquidarán exclusivamente las HH aprobadas hasta la fecha de corte (fecha de trabajo y aprobación), con la tarifa estándar actual del trabajador. Las liquidaciones cerradas conservan su importe histórico.</p><div class="grid"><a class="btn secondary" href="{{url_for('export_excel',month=month)}}">Exportar Excel</a><a class="btn secondary" href="{{url_for('export_pdf',month=month)}}">Exportar PDF</a></div></div><div class="card scroll"><h2>HH listas para liquidar</h2><table><tr><th>Trabajador</th><th>Registros</th><th>HH aprobadas</th><th>Acción</th></tr>{% for a in available %}<tr><td>{{a.name}}</td><td>{{a.n}}</td><td>{{a.q/4}}</td><td><form method="post"><input type="hidden" name="_csrf" value="{{csrf()}}"><input type="hidden" name="month" value="{{month}}"><input type="hidden" name="cutoff" value="{{cutoff}}"><input type="hidden" name="uid" value="{{a.id}}"><button>Liquidar</button></form></td></tr>{% endfor %}</table></div><div class="card scroll"><h2>Liquidaciones cerradas del periodo</h2><table><tr><th>ID</th><th>Trabajador</th><th>Periodo</th><th>Corte</th><th>HH</th><th>Total preliminar</th><th>Detalle</th></tr>{% for r in rows %}<tr><td>{{r.id}}</td><td>{{r.person}}</td><td>{{r.period}}</td><td>{{r.cutoff_date or 'Histórica'}}</td><td>{{r.quarters/4}}</td><td>{{fmt(r.cents)}}</td><td><a href="{{url_for('liquidation_detail',lid=r.id)}}">Ver</a></td></tr>{% endfor %}</table></div>''',month=month,cutoff=cutoff,available=available,rows=rows)
@app.route('/liquidations/<int:lid>')
@roles('admin')
def liquidation_detail(lid):
    with conn() as c:
        l=c.execute('SELECT l.*,u.name person FROM liquidations l JOIN users u ON u.id=l.user_id WHERE l.id=?',(lid,)).fetchone()
        if not l:abort(404)
        lines=c.execute('''SELECT d.*,e.work_date,p.code pc,a.code ac,e.description FROM liquidation_lines d JOIN entries e ON e.id=d.entry_id JOIN activities a ON a.id=e.activity_id JOIN projects p ON p.id=a.project_id WHERE d.liquidation_id=? ORDER BY e.work_date''',(lid,)).fetchall()
    return page('''<h1>Liquidación #{{l.id}}</h1><div class="card"><h2>{{l.person}} · {{l.period}}</h2><p><strong>{{l.quarters/4}} HH — {{fmt(l.cents)}}</strong></p><small>Importe preliminar sin impuestos ni retenciones.</small><div class="scroll"><table><tr><th>Fecha</th><th>Proyecto</th><th>Actividad</th><th>HH</th><th>Tarifa</th><th>Importe</th></tr>{% for r in lines %}<tr><td>{{r.work_date}}</td><td>{{r.pc}}</td><td>{{r.ac}}</td><td>{{r.quarters/4}}</td><td>{{fmt(r.rate_cents)}}</td><td>{{fmt(r.amount_cents)}}</td></tr>{% endfor %}</table></div></div>''',l=l,lines=lines)

def export_rows(month):
    with conn() as c:
        return c.execute('SELECT l.id,u.name,l.period,l.quarters,l.cents,l.created_at FROM liquidations l JOIN users u ON u.id=l.user_id WHERE l.period=? ORDER BY u.name',(month,)).fetchall()
@app.route('/export/xlsx')
@roles('admin')
def export_excel():
    month=request.args.get('month',period());date.fromisoformat(month+'-01')
    wb=Workbook();ws=wb.active;ws.title='Liquidaciones';ws.append(['ID','Trabajador','Periodo','HH','Total preliminar (S/)','Fecha cierre'])
    for r in export_rows(month):ws.append([r['id'],r['name'],r['period'],r['quarters']/4,r['cents']/100,r['created_at']])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='143A4E')
    for col,width in {'A':9,'B':32,'C':15,'D':13,'E':24,'F':25}.items():ws.column_dimensions[col].width=width
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    return send_file(buf,as_attachment=True,download_name=f'ALTIVIA_Liquidaciones_{month}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
@app.route('/export/pdf')
@roles('admin')
def export_pdf():
    month=request.args.get('month',period());date.fromisoformat(month+'-01')
    buf=io.BytesIO();pdf=canvas.Canvas(buf,pagesize=A4);y=790
    pdf.setFont('Helvetica-Bold',15);pdf.drawString(42,y,'ALTIVIA | Liquidaciones '+month);y-=32
    pdf.setFont('Helvetica',9)
    rows=export_rows(month)
    for r in rows:
        if y<65:pdf.showPage();y=790;pdf.setFont('Helvetica',9)
        pdf.drawString(42,y,f"#{r['id']}  {r['name'][:34]}  |  {r['quarters']/4:g} HH  |  {fmt(r['cents'])}");y-=19
    pdf.setFont('Helvetica-Oblique',8);pdf.drawString(42,40,'Importes preliminares; revisar obligaciones tributarias y contractuales.');pdf.save();buf.seek(0)
    return send_file(buf,as_attachment=True,download_name=f'ALTIVIA_Liquidaciones_{month}.pdf',mimetype='application/pdf')
@app.route('/import-matrix',methods=['GET','POST'])
@roles('admin')
def import_matrix():
    with conn() as c:
        projects=c.execute('SELECT * FROM projects ORDER BY code').fetchall()
        if request.method=='POST':
            import openpyxl,unicodedata
            from werkzeug.utils import secure_filename
            f=request.files.get('file');pid=int(request.form.get('project_id','0'))
            if not c.execute('SELECT 1 FROM projects WHERE id=?',(pid,)).fetchone():abort(400)
            if not f or not f.filename.lower().endswith('.xlsx'):raise ValueError('Sube archivo .xlsx')
            wb=openpyxl.load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True)
            if 'Estimado de Costos' not in wb.sheetnames:abort(400,'No se encuentra la hoja Estimado de Costos')
            ws=wb['Estimado de Costos']
            # Original template: B task code, C task, J:T professional, row 5 names, row 4 disciplines/roles.
            names=[ws.cell(5,k).value for k in range(10,21)]
            roles_excel=[ws.cell(4,k).value for k in range(10,21)]
            mapped={};created=[];existing=[]
            for idx,full in enumerate(names,10):
                if not full or norm(full) in ('tbd','por definir',''):continue
                people=c.execute('SELECT id,name FROM users').fetchall()
                exact=[u for u in people if norm(u['name'])==norm(full)]
                if len(exact)>1:raise ValueError('Nombre duplicado en usuarios: '+str(full))
                if exact:
                    mapped[idx]=exact[0]['id'];existing.append(str(full));continue
                login=unique_login(c,str(full))
                password=seed_password(str(full))
                uid=c.execute('INSERT INTO users(username,name,pw,role,must_change_pw) VALUES(?,?,?,?,1)',(login,str(full).strip(),generate_password_hash(password),'trabajador')).lastrowid
                mapped[idx]=uid;created.append(f'{full} (usuario: {login})')
                audit(c,'importar_usuario','user',uid,'rol trabajador; clave temporal')
            counter=0;rights=0;skipped=0
            for row in ws.iter_rows(min_row=10,values_only=True):
                code=row[1] if len(row)>2 else None;title=row[2] if len(row)>2 else None
                if isinstance(code,bool) or not isinstance(code,(int,float)) or not title:continue
                code=str(int(code)) if float(code).is_integer() else str(code)
                if code.endswith('000'):continue  # Summary lines, not billable activities
                allot=[]
                for col,uid in mapped.items():
                    value=row[col-1] if len(row)>=col else None
                    if isinstance(value,(int,float)) and not isinstance(value,bool) and value>0:
                        try:q=require_budget(str(value))
                        except (ValueError,InvalidOperation):continue
                        allot.append((uid,q))
                if not allot:skipped+=1;continue
                budget=sum(q for uid,q in allot)
                a=c.execute('SELECT id FROM activities WHERE project_id=? AND code=?',(pid,code)).fetchone()
                if a:
                    aid=a['id']
                    finalized=c.execute("SELECT 1 FROM entries WHERE activity_id=? AND state='liquidado' LIMIT 1",(aid,)).fetchone()
                    if finalized:raise ValueError('La matriz pretende modificar una actividad con HH liquidadas: '+code)
                    c.execute('UPDATE activities SET name=?,budget_quarters=?,active=1 WHERE id=?',(str(title).strip(),budget,aid))
                else:aid=c.execute('INSERT INTO activities(project_id,code,name,budget_quarters) VALUES(?,?,?,?)',(pid,code,str(title).strip(),budget)).lastrowid
                c.execute('DELETE FROM activity_assignments WHERE activity_id=?',(aid,))
                for uid,q in allot:
                    c.execute('INSERT INTO assignments(user_id,project_id) VALUES(?,?) ON CONFLICT DO NOTHING',(uid,pid))
                    c.execute('INSERT INTO activity_assignments(user_id,activity_id,planned_quarters) VALUES(?,?,?)',(uid,aid,q));rights+=1
                counter+=1
            c.execute('INSERT INTO import_log(project_id,filename,processed_at,actor_id,notes) VALUES(?,?,?,?,?)',(pid,secure_filename(f.filename),now(),user()['id'],f'{counter} actividades, {rights} asignaciones, {len(created)} usuarios nuevos'))
            audit(c,'importar_matriz','project',pid,f'{counter} actividades, {rights} cruces autorizados');c.commit()
            flash(f'Matriz importada: {counter} actividades, {rights} autorizaciones, {len(created)} usuarios nuevos. {skipped} filas ignoradas. Nuevos usuarios: '+(', '.join(created) if created else 'ninguno')+'.','ok')
            return redirect(url_for('admin'))
    return page("""<h1>Importar EDT/WBS y responsables desde Excel</h1><div class='card'><p>Plantilla ALTIVIA: hoja <strong>Estimado de Costos</strong>, códigos en B, actividades en C, nombres de los profesionales en fila 5 (J:T) y HH presupuestadas en cada cruce.</p><p class='warning'>Solo se importan cruces numéricos positivos. Se ignoran las celdas vacías, ceros y profesionales TBD. Los profesionales nuevos se crean automáticamente con rol Trabajador. Usuario inicial: último apellido disponible (con número si se repite). Clave temporal: primer nombre; deberán cambiarla al ingresar. Los nombres ya existentes se reutilizan sin cambiar sus roles. La importación actualiza el presupuesto y la matriz del proyecto elegido.</p><form method='post' enctype='multipart/form-data'><input type='hidden' name='_csrf' value='{{csrf()}}'><label>Proyecto existente</label><select name='project_id' required>{% for p in projects %}<option value='{{p.id}}'>{{p.code}} — {{p.name}}</option>{% endfor %}</select><label>Libro Excel (.xlsx)</label><input type='file' name='file' accept='.xlsx' required><button>Importar matriz y asignaciones</button></form></div>""",projects=projects)

@app.route('/audit')
@roles('admin')
def audit_view():
    with conn() as c:rows=c.execute('SELECT a.*,u.name person FROM audit a LEFT JOIN users u ON u.id=a.actor_id ORDER BY a.id DESC LIMIT 200').fetchall()
    return page('''<div class="card scroll"><h1>Auditoría (últimos 200 eventos)</h1><table><tr><th>Fecha</th><th>Actor</th><th>Acción</th><th>Entidad</th><th>ID</th><th>Detalle</th></tr>{% for r in rows %}<tr><td>{{r.at}}</td><td>{{r.person}}</td><td>{{r.action}}</td><td>{{r.target}}</td><td>{{r.target_id}}</td><td>{{r.detail}}</td></tr>{% endfor %}</table></div>''',rows=rows)
@app.errorhandler(403)
def forbidden(e):return page('<div class="card"><h1>Acceso denegado</h1><p>No tienes permisos para esta operación.</p></div>'),403
if __name__=='__main__':
    print('ALTIVIA HH CONTROL: http://127.0.0.1:8765')
    serve(app,host='127.0.0.1',port=8765,threads=4)
