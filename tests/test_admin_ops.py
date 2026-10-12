from pathlib import Path
import pytest
from streamlit.testing.v1 import AppTest
from werkzeug.security import check_password_hash
import db_sharepoint
from test_sharepoint import setup
from admin_ops import create_collaborator, reset_password, replace_assignments, set_project_status, name_parts


def test_first_name_first_surname_credentials_and_duplicates(setup):
    _, s, e=setup
    with e.begin() as c:
        first=create_collaborator(c,1,'José Antonio','Díaz','López')
        second=create_collaborator(c,1,'José Carlos','Paz','Ríos')
    users=s.snapshot()['data']['users']
    assert first['username']=='jose' and second['username']=='jose2'
    assert first['initial_password']=='diaz' and second['initial_password']=='paz'
    assert check_password_hash(users[str(first['id'])]['pw'],'diaz')
    assert users[str(first['id'])]['must_change_pw']==1
    assert 'diaz' not in str(s.snapshot()['data']['audit']).lower()
    assert name_parts('José Antonio Díaz López')['Primer_apellido']=='Díaz'


def test_reset_password_invalidates_live_session_and_requires_change(setup,monkeypatch):
    _, s, e=setup
    monkeypatch.setattr(db_sharepoint,'_engine',e)
    monkeypatch.setenv('DATA_BACKEND','sharepoint')
    worker=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    worker.session_state['uid']=3
    worker.run()
    with e.begin() as c:reset_password(c,1,3,'nuevo6')
    worker.run()
    assert not worker.exception and not worker.sidebar.radio
    assert any(b.label=='Ingresar' for b in worker.button)
    next(t for t in worker.text_input if t.label=='Usuario').set_value('worker')
    next(t for t in worker.text_input if t.label=='Contraseña').set_value('nuevo6')
    next(b for b in worker.button if b.label=='Ingresar').click().run()
    assert any(t.label=='Nueva contraseña (mínimo 6 caracteres)' for t in worker.text_input)
    next(t for t in worker.text_input if t.label=='Nueva contraseña (mínimo 6 caracteres)').set_value('propia6')
    next(t for t in worker.text_input if t.label=='Repetir nueva contraseña').set_value('propia6')
    next(b for b in worker.button if b.label=='Actualizar contraseña').click().run()
    assert not worker.exception and worker.sidebar.radio[0].value=='Registrar horas'
    assert check_password_hash(s.snapshot()['data']['users']['3']['pw'],'propia6')
    assert s.snapshot()['data']['users']['3']['must_change_pw']==0


def test_assignment_edits_preserve_planned_hours_and_history(setup):
    _, s, e=setup
    with e.begin() as c:
        c.execute('INSERT INTO assignments(user_id,project_id) VALUES(3,1)')
        c.execute('INSERT INTO activity_assignments(user_id,activity_id,planned_quarters) VALUES(3,1,77)')
        c.execute("INSERT INTO activities(project_id,code,name,budget_quarters) VALUES(1,'B','Second activity',100)")
        c.execute("INSERT INTO entries(user_id,activity_id,work_date,quarters,description,state,submission_key) VALUES(3,1,'2026-10-10',4,'Historical hours','aprobado','historic')")
    with e.begin() as c:replace_assignments(c,1,3,1,[1,2],[1])
    data=s.snapshot()['data']
    assert next(r for r in data['activity_assignments'].values() if r['activity_id']==1)['planned_quarters']==77
    with e.begin() as c:replace_assignments(c,1,3,1,[2],[1,2])
    data=s.snapshot()['data']
    assert {r['activity_id'] for r in data['activity_assignments'].values()}=={2}
    assert data['entries']['1']['activity_id']==1 and data['entries']['1']['state']=='aprobado'
    assert data['activities']['1']['budget_quarters']==400
    with pytest.raises(ValueError,match='cambiaron'):
        with e.begin() as c:replace_assignments(c,1,3,1,[1],[1,2])
    with e.begin() as c:replace_assignments(c,1,3,1,[],[2],False)
    assert not s.snapshot()['data']['assignments'] and not s.snapshot()['data']['activity_assignments']
    assert s.snapshot()['data']['entries']['1']['quarters']==4


def test_admin_mutations_reject_worker_and_cross_project_activity(setup):
    _, _, e=setup
    for operation in [lambda c:reset_password(c,3,2,'nuevo6'),lambda c:create_collaborator(c,3,'Juan','Paz'),lambda c:replace_assignments(c,3,3,1,[1],[]),lambda c:set_project_status(c,3,1,'terminado','activo')]:
        with pytest.raises(ValueError,match='administrador'):
            with e.begin() as c:operation(c)
    with e.begin() as c:
        c.execute("INSERT INTO projects(code,name) VALUES('OTHER','Other project')")
        c.execute("INSERT INTO activities(project_id,code,name) VALUES(2,'X','Other activity')")
    with pytest.raises(ValueError,match='no pertenece'):
        with e.begin() as c:replace_assignments(c,1,3,1,[2],[])


def test_finish_project_blocks_new_hours_but_preserves_history(setup,monkeypatch):
    _, s, e=setup
    monkeypatch.setattr(db_sharepoint,'_engine',e)
    monkeypatch.setenv('DATA_BACKEND','sharepoint')
    with e.begin() as c:
        c.execute('INSERT INTO assignments(user_id,project_id) VALUES(3,1)')
        c.execute('INSERT INTO activity_assignments(user_id,activity_id) VALUES(3,1)')
        c.execute("INSERT INTO entries(user_id,activity_id,work_date,quarters,description,state,submission_key) VALUES(3,1,'2026-10-10',4,'Past work','aprobado','past')")
    admin=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    admin.session_state['uid']=1
    admin.run().sidebar.radio[0].set_value('Proyectos y actividades').run()
    next(b for b in admin.button if b.label=='Terminar proyecto').click().run()
    assert not admin.exception and s.snapshot()['data']['projects']['1']['status']=='terminado'
    worker=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    worker.session_state['uid']=3
    worker.run()
    assert not any(b.label=='Enviar para aprobación' for b in worker.button)
    worker.sidebar.radio[0].set_value('Mi historial').run()
    assert worker.dataframe[0].value.iloc[0]['Descripción']=='Past work'


def test_admin_forms_create_reset_and_edit_authorizations(setup,monkeypatch):
    _, s, e=setup
    monkeypatch.setattr(db_sharepoint,'_engine',e)
    monkeypatch.setenv('DATA_BACKEND','sharepoint')
    with e.begin() as c:
        c.execute('INSERT INTO assignments(user_id,project_id) VALUES(3,1)')
        c.execute('INSERT INTO activity_assignments(user_id,activity_id,planned_quarters) VALUES(3,1,40)')
    admin=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    admin.session_state['uid']=1
    admin.run().sidebar.radio[0].set_value('Usuarios y tarifas').run()
    next(t for t in admin.text_input if t.label=='Nombres').set_value('Pedro Luis')
    next(t for t in admin.text_input if t.label=='Primer apellido').set_value('Paz')
    next(b for b in admin.button if b.label=='Crear colaborador').click().run()
    assert not admin.exception and not admin.error
    created=next(r for r in s.snapshot()['data']['users'].values() if r['username']=='pedro')
    assert check_password_hash(created['pw'],'paz') and created['must_change_pw']==1
    next(t for t in admin.selectbox if t.label=='Colaborador a restablecer').set_value(3)
    next(t for t in admin.text_input if t.label=='Nueva contraseña temporal (mínimo 6 caracteres)').set_value('reset6')
    next(t for t in admin.text_input if t.label=='Repetir contraseña temporal').set_value('reset6')
    next(b for b in admin.button if b.label=='Restablecer contraseña').click().run()
    assert not admin.exception and not admin.error
    assert check_password_hash(s.snapshot()['data']['users']['3']['pw'],'reset6')
    next(t for t in admin.selectbox if t.label=='Colaborador').set_value('3 · worker').run()
    assert admin.multiselect[0].value==['A · Test']
    admin.multiselect[0].set_value([])
    next(b for b in admin.button if b.label=='Guardar asignaciones del colaborador').click().run()
    assert not admin.exception and not admin.error
    assert not s.snapshot()['data']['activity_assignments']
    assert s.snapshot()['data']['assignments']


def test_changed_project_access_cannot_be_overwritten_from_stale_form(setup):
    _, _, e=setup
    with pytest.raises(ValueError,match='acceso al proyecto cambió'):
        with e.begin() as c:replace_assignments(c,1,3,1,[],[],True,True)


def test_excel_import_uses_reviewed_name_and_keeps_existing_credentials(setup,monkeypatch):
    import io
    import streamlit as st
    from openpyxl import Workbook
    _, s, e=setup
    monkeypatch.setattr(db_sharepoint,'_engine',e)
    monkeypatch.setenv('DATA_BACKEND','sharepoint')
    workbook=Workbook();sheet=workbook.active;sheet.title='Estimado de Costos'
    sheet.cell(5,10,'Carlos Mendoza Romero')
    sheet.cell(10,2,10110);sheet.cell(10,3,'Imported work');sheet.cell(10,10,2.5)
    payload=io.BytesIO();workbook.save(payload)
    class Upload:
        name='synthetic_budget.xlsx'
        def getvalue(self):return payload.getvalue()
    monkeypatch.setattr(st,'file_uploader',lambda *args,**kwargs:Upload())
    admin=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    admin.session_state['uid']=1
    admin.run().sidebar.radio[0].set_value('Importar Excel').run()
    assert not admin.exception
    next(b for b in admin.button if b.label=='Confirmar importación de matriz').click().run()
    assert not admin.exception and not admin.error
    data=s.snapshot()['data'];created=next(r for r in data['users'].values() if r['username']=='carlos')
    assert check_password_hash(created['pw'],'mendoza') and created['must_change_pw']==1
    assert created['name']=='Carlos Mendoza Romero'
    activity=next(r for r in data['activities'].values() if r['code']=='10110')
    assert activity['budget_quarters']==10
    hashed=created['pw'];count=len(data['users'])
    next(b for b in admin.button if b.label=='Confirmar importación de matriz').click().run()
    assert not admin.exception and not admin.error
    data=s.snapshot()['data']
    assert len(data['users'])==count and data['users'][str(created['id'])]['pw']==hashed
