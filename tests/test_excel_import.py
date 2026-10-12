import io
import pytest
from openpyxl import Workbook
from excel_import import read_matrix,resolve_review,vertical_total
from test_sharepoint import setup


def payload(sheet):
    stream=io.BytesIO();sheet.parent.save(stream);return stream.getvalue()


def matrix():
    book=Workbook();sheet=book.active;sheet.title='Estimado de Costos'
    sheet.cell(5,10,'Carlos Mendoza Romero')
    return sheet


def test_unbounded_columns_rows_and_duplicate_professional_hours():
    s=matrix();s.cell(5,34,'Eva Usnayo Perez');s.cell(5,65,'Carlos Mendoza Romero');s.cell(5,80,'Ana Paz Ramos')
    s.cell(220,2,'90910');s.cell(220,3,'Activity beyond old limits')
    for col,hours in ((10,1),(34,2),(65,3),(80,4)):s.cell(220,col,hours)
    m=read_matrix(payload(s));r=m['activities'][0]
    assert not m['errors'] and m['last_column']==80 and m['last_row']==220
    assert dict(r['cross'])=={'Carlos Mendoza Romero':16,'Eva Usnayo Perez':8,'Ana Paz Ramos':16}
    assert r['budget']==40


def test_general_totals_blank_separators_and_leaf_ending_in_00():
    s=matrix()
    for row,code,title,hours in [(10,20000,'General summary',7),(11,20100,'Actual task',3),(12,20200,'Actual task 2',4),(14,30000,'Another summary',2),(15,30100,'Actual task 3',2),(17,'HORAS TOTALES','',9),(18,'COSTO TOTAL','',900)]:
        s.cell(row,2,code);s.cell(row,3,title);s.cell(row,10,hours)
    m=read_matrix(payload(s))
    assert [r['code'] for r in m['activities']]==['20100','20200','30100']
    assert len(m['excluded'])==4 and m['source_total']==9
    assert sum(r['budget'] for r in m['activities'])==36


def test_unassigned_positions_preserve_budget_and_no_edt_requires_review():
    s=matrix();s.cell(5,34,'TBD');s.cell(4,34,'Environmental engineer')
    s.cell(10,2,10110);s.cell(10,3,'Task');s.cell(10,10,2);s.cell(10,34,3);s.cell(10,36,4)
    s.cell(12,3,'Uncoded work');s.cell(12,10,1)
    m=read_matrix(payload(s));r=m['activities'][0]
    assert r['budget']==36 and r['unassigned']==28 and len(m['review'])==1
    assert len(m['positions'])==3 and len(m['persons'])==1
    with pytest.raises(ValueError,match='EDT'):
        resolve_review(m,[{'Incluir':True,'EDT':''}])
    selected=resolve_review(m,[{'Incluir':True,'EDT':'10120'}])
    assert len(selected)==2 and selected[1]['budget']==4
    assert len(resolve_review(m,[{'Incluir':False,'EDT':''}]))==1


def test_invalid_hours_and_uncached_formulas_are_errors_not_silent_skips():
    s=matrix()
    for row,value in [(10,.3),(11,-1),(12,'eight'),(13,'=2*4')]:
        s.cell(row,2,10100+row);s.cell(row,3,'Task');s.cell(row,10,value)
    m=read_matrix(payload(s))
    assert len(m['errors'])==4 and any('fórmula' in e for e in m['errors'])


def test_vertical_formula_summaries_and_no_false_horizontal_summary():
    assert vertical_total('=SUM(J11:J13)',10)
    assert not vertical_total('=SUM(J10:AH10)',10)
    s=matrix();s.cell(10,2,10100);s.cell(10,3,'Summary');s.cell(10,10,'=SUM(J11:J12)')
    for row in (11,12):
        s.cell(row,2,10100+row);s.cell(row,3,'Task');s.cell(row,10,2)
    m=read_matrix(payload(s));assert len(m['activities'])==2 and not m['errors']
    assert m['excluded'][0]['reason']=='Resumen general / suma de subtareas'


def test_duplicate_edt_codes_block_confirmation():
    s=matrix()
    for row in (10,11):
        s.cell(row,2,10110);s.cell(row,3,'Task');s.cell(row,10,2)
    assert any('repetidos' in error for error in read_matrix(payload(s))['errors'])


def test_final_leaf_ending_in_000_is_not_lost_and_blank_row_does_not_hide_work():
    s=matrix()
    for row,code,title in [(10,90000,'Summary'),(11,90100,'Leaf'),(12,91000,'Final deliverable'),(14,91100,'More work after blank row')]:
        s.cell(row,2,code);s.cell(row,3,title);s.cell(row,10,2)
    m=read_matrix(payload(s))
    assert [r['code'] for r in m['activities']]==['90100','91000','91100']


def test_import_screen_requires_review_then_maps_late_users_and_full_budget(setup,monkeypatch):
    from pathlib import Path
    import streamlit as st
    import db_sharepoint
    from streamlit.testing.v1 import AppTest
    _,store,engine=setup
    monkeypatch.setattr(db_sharepoint,'_engine',engine)
    monkeypatch.setenv('DATA_BACKEND','sharepoint')
    s=matrix();s.cell(5,10,'worker');s.cell(5,65,'Ana Paz Ramos');s.cell(5,80,'TBD')
    s.cell(220,2,90910);s.cell(220,3,'Late activity')
    s.cell(220,10,2);s.cell(220,65,3);s.cell(220,80,4)
    s.cell(222,3,'Uncoded heading');s.cell(222,10,1)
    blob=payload(s)
    class Upload:
        name='synthetic_extended.xlsx'
        def getvalue(self):return blob
    monkeypatch.setattr(st,'file_uploader',lambda *args,**kwargs:Upload())
    app=AppTest.from_file(Path('app.py').resolve(),default_timeout=20)
    app.session_state['uid']=1
    app.run().sidebar.radio[0].set_value('Importar Excel').run()
    assert not app.exception and not app.error
    assert next(b for b in app.button if b.label=='Confirmar importación de matriz').disabled
    next(c for c in app.checkbox if c.label=='He revisado las filas excluidas y las horas sin colaborador asignado').check().run()
    next(b for b in app.button if b.label=='Confirmar importación de matriz').click().run()
    assert not app.exception and not app.error
    data=store.snapshot()['data']
    ana=next(r for r in data['users'].values() if r['username']=='ana')
    activity=next(r for r in data['activities'].values() if r['code']=='90910')
    assert activity['budget_quarters']==36
    assert not any(r['name']=='TBD' for r in data['users'].values())
    assert {(r['user_id'],r['planned_quarters']) for r in data['activity_assignments'].values() if r['activity_id']==activity['id']}=={(3,8),(ana['id'],12)}
    assert not any(r['name']=='Uncoded heading' for r in data['activities'].values())
